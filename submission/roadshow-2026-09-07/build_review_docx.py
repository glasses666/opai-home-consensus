"""Build the review document without changing the original submission.

Run with the Python runtime returned by load_workspace_dependencies.
"""
from pathlib import Path
import re

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "01-参赛方案优化稿.md"
OUTPUT = ROOT / "01-参赛方案优化稿.docx"
FONT = "Arial Unicode MS"


def inline(paragraph, text):
    for piece in re.split(r"(\*\*.*?\*\*|\[[^\]]+\]\(https?://[^)]+\))", text):
        link = re.fullmatch(r"\[([^\]]+)\]\((https?://[^)]+)\)", piece)
        if link:
            label, url = link.groups()
            rel = paragraph.part.relate_to(
                url,
                "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                is_external=True,
            )
            node = OxmlElement("w:hyperlink")
            node.set(qn("r:id"), rel)
            run = OxmlElement("w:r")
            props = OxmlElement("w:rPr")
            color = OxmlElement("w:color")
            color.set(qn("w:val"), "245477")
            props.append(color)
            run.append(props)
            text_node = OxmlElement("w:t")
            text_node.text = label
            run.append(text_node)
            node.append(run)
            paragraph._p.append(node)
        else:
            bold = piece.startswith("**") and piece.endswith("**")
            run = paragraph.add_run(piece[2:-2] if bold else piece)
            run.bold = bold


def table(doc, rows):
    tbl = doc.add_table(rows=len(rows), cols=len(rows[0]))
    tbl.autofit = False
    widths = [1.25, 2.85, 2.9] if len(rows[0]) == 3 else [7 / len(rows[0])] * len(rows[0])
    for col, width in zip(tbl.columns, widths):
        col.width = Inches(width)
    for idx, values in enumerate(rows):
        for cell, value, width in zip(tbl.rows[idx].cells, values, widths):
            cell.width = Inches(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tc_pr = cell._tc.get_or_add_tcPr()
            borders = OxmlElement("w:tcBorders")
            for edge in ("top", "left", "bottom", "right"):
                border = OxmlElement(f"w:{edge}")
                for name, val in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
                    border.set(qn(f"w:{name}"), val)
                borders.append(border)
            tc_pr.append(borders)
            margins = OxmlElement("w:tcMar")
            for edge in ("top", "left", "bottom", "right"):
                node = OxmlElement(f"w:{edge}")
                node.set(qn("w:w"), "95")
                node.set(qn("w:type"), "dxa")
                margins.append(node)
            tc_pr.append(margins)
            shading = OxmlElement("w:shd")
            shading.set(qn("w:fill"), "243F50" if idx == 0 else ("F1F4F6" if idx % 2 == 0 else "FFFFFF"))
            tc_pr.append(shading)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = Pt(16)
            inline(p, value)
            for run in p.runs:
                run.font.size = Pt(10.5)
                if idx == 0:
                    run.bold = True
                    run.font.color.rgb = RGBColor(255, 255, 255)
        row_pr = tbl.rows[idx]._tr.get_or_add_trPr()
        row_pr.append(OxmlElement("w:cantSplit"))
        if idx == 0:
            row_pr.append(OxmlElement("w:tblHeader"))
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def build():
    doc = Document()
    for border in doc.styles.element.xpath('.//w:pBdr'):
        border.getparent().remove(border)
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.65)
    section.left_margin = section.right_margin = Inches(0.75)
    section.footer_distance = Inches(0.3)
    for grid in section._sectPr.xpath('./w:docGrid'):
        grid.getparent().remove(grid)
    for name, size in (("Normal", 11), ("Title", 25), ("Subtitle", 13), ("Heading 1", 17), ("Heading 2", 13)):
        style = doc.styles[name]
        style.font.name = FONT
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        fonts = style._element.get_or_add_rPr().get_or_add_rFonts()
        for key in list(fonts.attrib):
            del fonts.attrib[key]
        for key in ("ascii", "hAnsi", "eastAsia", "cs"):
            fonts.set(qn(f"w:{key}"), FONT)
        style.paragraph_format.space_after = Pt(7)
        style.paragraph_format.line_spacing = Pt(size + 6)
        if name.startswith("Heading"):
            style.font.bold = True
            style.paragraph_format.space_before = Pt(11)
            style.paragraph_format.keep_with_next = True
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    lines = SOURCE.read_text().splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if line == "<!-- pagebreak -->":
            doc.add_page_break()
        elif line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                raw = lines[i].strip()
                if not re.fullmatch(r"[|\s:-]+", raw):
                    rows.append([part.strip() for part in raw.strip("|").split("|")])
                i += 1
            table(doc, rows)
            continue
        elif line.startswith("# "):
            doc.add_paragraph(line[2:], "Title")
        elif line.startswith("## "):
            doc.add_paragraph(line[3:], "Heading 1")
        elif line.startswith("### "):
            doc.add_paragraph(line[4:], "Heading 2")
        elif line:
            paragraph = doc.add_paragraph()
            if line.startswith("[1]") or line.startswith("[2]") or line.startswith("[3]") or line.startswith("[4]") or line.startswith("本版结构参考"):
                paragraph.paragraph_format.line_spacing = Pt(13)
                paragraph.paragraph_format.space_after = Pt(4)
            inline(paragraph, line)
            if line.startswith("[1]") or line.startswith("[2]") or line.startswith("[3]") or line.startswith("[4]") or line.startswith("本版结构参考"):
                for run in paragraph.runs:
                    run.font.size = Pt(9)
        i += 1
    doc.core_properties.title = "元界视创欧派 AI 家装共识工作台"
    doc.core_properties.subject = "区域路演方案优化"
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
